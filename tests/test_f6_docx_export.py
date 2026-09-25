from __future__ import annotations

from io import BytesIO

from docx import Document

from src.search_intelligence.f6_docx_export import (
    build_editable_companion_docx,
    build_fillable_starter_docx,
)


def _document_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def test_fillable_starter_is_local_editable_word_document() -> None:
    content = build_fillable_starter_docx()

    assert content.startswith(b"PK")
    text = _document_text(content)
    assert "[Vorname Nachname]" in text
    assert "Bewerbung als [Stellenbezeichnung]" in text
    assert "[Kurzprofil:" in text
    assert "[Projekt 1]" in text


def test_editable_companion_uses_final_f6_text_values() -> None:
    values = {
        "base_application_letter": {
            "header.name": "Jens Haberle",
            "header.tagline": "Data Engineering",
            "header.contact": "Hannover · jens@example.test",
            "recipient.block": "Eraneos\nz. Hd. Paulina Krzeminski",
            "date": "25.09.2026",
            "subject": "Bewerbung als Data Engineer",
            "salutation": "Sehr geehrte Frau Krzeminski,",
            "body.paragraph_1": "Ein vollständiger, stellenbezogener erster Absatz.",
            "body.paragraph_2": "Ein zweiter vollständiger Absatz mit Engineering-Bezug.",
            "closing.formula": "Mit freundlichen Grüßen",
            "closing.name": "Jens Haberle",
        },
        "base_cv": {
            "p1.header.name": "Jens Haberle",
            "p1.header.tagline": "Data Engineering",
            "p1.header.contact": "Hannover · jens@example.test",
            "p1.short_profile": "Gezieltes Data-Engineering-Kurzprofil.",
            "p1.competency_profile": "Python · PostgreSQL · Data Quality",
            "p1.experience.1": "CARIAD SE · System Development Engineer",
            "p2.education.1": "TU Braunschweig · Diplom-Wirtschaftsingenieur",
            "p2.project.1": "Pedestrian Dataset Engineering",
            "p2.skills.data_ai": "Python\nSQL / PostgreSQL",
            "p2.skills.engineering": "System Engineering",
            "p2.skills.tools": "GitHub Actions",
            "p2.skills.languages": "Deutsch\nEnglisch",
            "p2.footer.date": "Hannover, 25. September 2026",
            "p2.footer.name": "Jens Haberle",
        },
    }

    content = build_editable_companion_docx(
        values_by_document=values,
        company_name="Eraneos",
        title="Data Engineer",
    )

    text = _document_text(content)
    assert "Eraneos" in text
    assert "Sehr geehrte Frau Krzeminski," in text
    assert "Gezieltes Data-Engineering-Kurzprofil." in text
    assert "Hannover, 25. September 2026" in text



def test_editable_companion_collapses_exact_duplicate_extraction_lines() -> None:
    values = {
        "base_application_letter": {
            "header.name": "Jens Haberle",
            "recipient.block": "Eraneos",
            "date": "25.09.2026",
            "subject": "Bewerbung als Data Engineer",
            "salutation": "Guten Tag,",
            "body.paragraph_1": "Kurzer Absatz.",
            "closing.formula": "Freundliche Grüße",
            "closing.name": "Jens Haberle",
        },
        "base_cv": {
            "p1.header.name": "Jens Haberle",
            "p1.short_profile": "Kurzprofil",
            "p1.competency_profile": "Python",
            "p1.experience.5": (
                "Alstom Transport Deutschland GmbH · Praktikant / Diplomand\n"
                "Diplomarbeit zu technischen Alternativen.\n"
                "Diplomarbeit zu technischen Alternativen.\n"
                "Technische Dokumentation.\n"
                "Technische Dokumentation."
            ),
            "p2.footer.name": "Jens Haberle",
        },
    }

    content = build_editable_companion_docx(
        values_by_document=values,
        company_name="Eraneos",
        title="Data Engineer",
    )

    text = _document_text(content)
    assert text.count("Diplomarbeit zu technischen Alternativen.") == 1
    assert text.count("Technische Dokumentation.") == 1

    document = Document(BytesIO(content))
    assert len(document.sections) == 2
    for section in document.sections:
        assert round(section.page_width.mm) == 210
        assert round(section.page_height.mm) == 297
