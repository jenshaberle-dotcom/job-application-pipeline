"""Local editable DOCX companions for F6.

PDF remains the canonical layout/pixel authority. DOCX output is an explicitly
editable companion generated from the same final text values; Word may reflow
fonts and spacing and therefore never carries F6 template authority.

The starter template is deliberately provider-free and contains placeholders
only. It gives users who cannot or do not want to upload private documents a
usable local starting point without inventing candidate facts.
"""
from __future__ import annotations

from io import BytesIO
from typing import Mapping

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor


ZoneValues = Mapping[str, Mapping[str, str]]


def _text(
    values: ZoneValues,
    document_type: str,
    zone_id: str,
    fallback: str = "",
) -> str:
    document = values.get(document_type)
    if not isinstance(document, Mapping):
        return fallback
    value = str(document.get(zone_id) or "").strip()
    if not value:
        return fallback

    # PDF extraction can expose duplicate logical lines when source text objects overlap.
    # The verified PDF remains untouched; only the editable Word companion is normalized.
    seen: set[str] = set()
    cleaned: list[str] = []
    for raw_line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        normalized = " ".join(raw_line.split())
        if normalized and normalized in seen:
            continue
        if normalized:
            seen.add(normalized)
        cleaned.append(raw_line.strip())
    compact = "\n".join(cleaned).strip()
    return compact or fallback


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.02

    title = document.styles["Title"]
    title.font.name = "Aptos"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x16, 0x2A, 0x45)

    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Aptos"
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x16, 0x2A, 0x45)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)
    document.styles["Heading 1"].font.size = Pt(12.5)
    document.styles["Heading 2"].font.size = Pt(10.5)


def _configure_letter_section(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(13)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)


def _add_identity_header(
    document: Document,
    *,
    name: str,
    tagline: str,
    contact: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(24)
    if tagline:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(tagline)
        run.bold = True
        run.font.size = Pt(10)
    if contact:
        paragraph = document.add_paragraph(contact)
        paragraph.paragraph_format.space_after = Pt(10)


def _add_letter(
    document: Document,
    values: ZoneValues,
    *,
    company_name: str,
    title: str,
) -> None:
    kind = "base_application_letter"
    _add_identity_header(
        document,
        name=_text(values, kind, "header.name", "[Name]"),
        tagline=_text(values, kind, "header.tagline", "[Profil / Schwerpunkt]"),
        contact=_text(values, kind, "header.contact", "[Kontakt]"),
    )

    meta = document.add_table(rows=1, cols=2)
    meta.autofit = True
    left, right = meta.rows[0].cells
    left.text = _text(values, kind, "recipient.block", company_name or "[Unternehmen]")
    right.text = _text(values, kind, "date", "[Datum]")
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()

    subject = document.add_paragraph()
    subject.paragraph_format.space_after = Pt(8)
    run = subject.add_run(
        _text(values, kind, "subject", f"Bewerbung als {title}" if title else "[Betreff]")
    )
    run.bold = True

    document.add_paragraph(_text(values, kind, "salutation", "[Anrede]"))
    for index in range(1, 7):
        value = _text(values, kind, f"body.paragraph_{index}")
        if value:
            paragraph = document.add_paragraph(value)
            paragraph.paragraph_format.space_after = Pt(8)

    document.add_paragraph(
        _text(values, kind, "closing.formula", "Mit freundlichen Grüßen")
    )
    # PDF signature art is deliberately not copied into the editable companion.
    # This keeps the companion editable and avoids duplicate/overlay signatures.
    document.add_paragraph()
    document.add_paragraph(_text(values, kind, "closing.name", "[Name]"))


def _add_cv(
    document: Document,
    values: ZoneValues,
) -> None:
    kind = "base_cv"
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)

    _add_identity_header(
        document,
        name=_text(values, kind, "p1.header.name", "[Name]"),
        tagline=_text(values, kind, "p1.header.tagline", "[Profil / Schwerpunkt]"),
        contact=_text(values, kind, "p1.header.contact", "[Kontakt]"),
    )

    document.add_heading("Kurzprofil", level=1)
    profile = document.add_table(rows=1, cols=2)
    profile.autofit = True
    profile.rows[0].cells[0].text = _text(
        values, kind, "p1.short_profile", "[Kurzprofil]"
    )
    profile.rows[0].cells[1].text = _text(
        values, kind, "p1.competency_profile", "[Kompetenzprofil]"
    )

    document.add_heading("Beruflicher Werdegang", level=1)
    for index in range(1, 6):
        value = _text(values, kind, f"p1.experience.{index}")
        if value:
            document.add_paragraph(value)

    document.add_page_break()
    document.add_heading("Ausbildung & Weiterbildung", level=1)
    for index in range(1, 4):
        value = _text(values, kind, f"p2.education.{index}")
        if value:
            document.add_paragraph(value)

    document.add_heading("Eigene Projekte", level=1)
    for index in range(1, 3):
        value = _text(values, kind, f"p2.project.{index}")
        if value:
            document.add_paragraph(value)

    document.add_heading("Kenntnisse", level=1)
    skills = document.add_table(rows=1, cols=4)
    skill_zones = (
        "p2.skills.data_ai",
        "p2.skills.engineering",
        "p2.skills.tools",
        "p2.skills.languages",
    )
    for cell, zone_id in zip(skills.rows[0].cells, skill_zones, strict=True):
        cell.text = _text(values, kind, zone_id, "[Kenntnisse]")

    footer_date = _text(values, kind, "p2.footer.date")
    footer_name = _text(values, kind, "p2.footer.name")
    if footer_date or footer_name:
        document.add_paragraph()
        if footer_date:
            document.add_paragraph(footer_date)
        if footer_name:
            document.add_paragraph(footer_name)


def build_editable_companion_docx(
    *,
    values_by_document: ZoneValues,
    company_name: str,
    title: str,
) -> bytes:
    """Build one editable Word companion from the final F6 text model."""

    document = Document()
    _configure_styles(document)
    _configure_letter_section(document)
    _add_letter(
        document,
        values_by_document,
        company_name=company_name,
        title=title,
    )
    _add_cv(document, values_by_document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _starter_values() -> dict[str, dict[str, str]]:
    return {
        "base_application_letter": {
            "header.name": "[Vorname Nachname]",
            "header.tagline": "[Berufsprofil / Schwerpunkte]",
            "header.contact": "[Adresse] · [Telefon] · [E-Mail]",
            "recipient.block": "[Unternehmen]\n[z. Hd. Ansprechpartner:in]\n[Adresse]",
            "date": "[Datum]",
            "subject": "Bewerbung als [Stellenbezeichnung]",
            "salutation": "[Anrede]",
            "body.paragraph_1": "[Warum diese Position und dieses Unternehmen?]",
            "body.paragraph_2": "[Relevante Berufserfahrung und belastbare Beispiele]",
            "body.paragraph_3": "[Relevante aktuelle Praxis / Projekte / Kompetenzen]",
            "body.paragraph_4": "[Welchen konkreten Mehrwert bringen Sie ein?]",
            "body.paragraph_5": "[Optional: Motivation / Transfer / Zusammenarbeit]",
            "body.paragraph_6": "[Kurzer Abschluss und Gesprächswunsch]",
            "closing.formula": "Mit freundlichen Grüßen",
            "closing.name": "[Vorname Nachname]",
        },
        "base_cv": {
            "p1.header.name": "[Vorname Nachname]",
            "p1.header.tagline": "[Berufsprofil / Schwerpunkte]",
            "p1.header.contact": "[Adresse] · [Telefon] · [E-Mail]",
            "p1.short_profile": "[Kurzprofil: 4–6 Sätze zu Erfahrung, Fokus und Mehrwert]",
            "p1.competency_profile": "[4–6 Kernkompetenzen]",
            "p1.experience.1": "[MM/JJJJ – heute] [Unternehmen] · [Rolle]\n• [Aufgabe / Ergebnis]\n• [Aufgabe / Ergebnis]",
            "p1.experience.2": "[MM/JJJJ – MM/JJJJ] [Unternehmen] · [Rolle]\n• [Aufgabe / Ergebnis]",
            "p1.experience.3": "[MM/JJJJ – MM/JJJJ] [Unternehmen] · [Rolle]\n• [Aufgabe / Ergebnis]",
            "p1.experience.4": "[Weitere Station]",
            "p1.experience.5": "[Weitere Station]",
            "p2.education.1": "[Ausbildung / Studium]",
            "p2.education.2": "[Weiterbildung]",
            "p2.education.3": "[Zertifikate]",
            "p2.project.1": "[Projekt 1]\n• [Technologien / Ergebnis]",
            "p2.project.2": "[Projekt 2]\n• [Technologien / Ergebnis]",
            "p2.skills.data_ai": "[Data / AI]",
            "p2.skills.engineering": "[Engineering]",
            "p2.skills.tools": "[Tools]",
            "p2.skills.languages": "[Sprachen]",
            "p2.footer.date": "[Ort, Datum]",
            "p2.footer.name": "[Vorname Nachname]",
        },
    }


def build_fillable_starter_docx() -> bytes:
    """Build a provider-free, fillable Word starter for users without source PDFs."""

    return build_editable_companion_docx(
        values_by_document=_starter_values(),
        company_name="[Unternehmen]",
        title="[Stellenbezeichnung]",
    )


__all__ = [
    "build_editable_companion_docx",
    "build_fillable_starter_docx",
]
